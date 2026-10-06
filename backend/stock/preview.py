"""公海历史成交表：按客户编码预览，确认前不落库。"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from typing import Any

from openpyxl import load_workbook

from ..storage import new_id
from . import catalog
from .storage import read_object

HEADER_FIELDS = {
    "customer_code": ("客户编码", "客户编号", "编码"),
    "company": ("客户名称", "公司全称", "公司名称", "企业名称"),
    "contact": ("联系人", "姓名"),
    "mobile": ("手机", "手机号"),
    "landline": ("座机", "固定电话"),
    "role": ("职务", "职位", "默认角色"),
    "product": ("历史成交产品", "成交产品", "已购经营范围", "已购产品", "已购"),
    "sales": ("销售", "所属销售"),
    "department": ("部门",),
    "coach": ("军师",),
    "email": ("邮箱", "电子邮箱"),
    "province": ("省份", "省"),
    "city": ("城市", "市"),
    "signals": ("公开信号", "已知信号"),
    "notes": ("备注", "说明"),
}

REQUIRED = ("customer_code", "company", "contact", "mobile")
WEAK_ROLES = {"hr", "经理", "老师", "先生", "女士", "培训"}
_SPLIT = re.compile(r"[、，,;；+\n/]+")
_MOBILE = re.compile(r"^1[3-9]\d{9}$")


def build_preview(content: bytes) -> dict[str, Any]:
    rows, headers = _read_sheet(content)
    mapping = _header_map(headers)
    missing = [field for field in REQUIRED if field not in mapping]
    if missing:
        return _blocked_file("不是存量导入模板，需要客户编码、客户名称、联系人、手机这几列")
    parsed = [_parse_row(raw, mapping, excel_row) for excel_row, raw in rows]
    parsed = [item for item in parsed if item["has_content"]]
    blockers: list[dict[str, str]] = []
    for item in parsed:
        if not item["customer_code"]:
            blockers.append(_issue("B1", item, "第 %s 行无客户编码，无法收成一家公司" % item["excel_row"]))
        elif not item["company"]:
            blockers.append(_issue("B2", item, "编码 %s 无客户名称" % item["customer_code"]))
    grouped = _group(parsed, blockers)
    warnings: list[dict[str, str]] = []
    hints: list[dict[str, str]] = []
    companies = [_company_card(code, items, warnings, hints) for code, items in grouped.items()]
    _cross_contacts(companies, warnings)
    _link_known_pairs(companies)
    _sales_columns(parsed, warnings)
    summary = {
        "excel_rows": len(parsed),
        "companies": len(companies),
        "contacts": sum(len(item["contacts"]) for item in companies),
        "multi_contact_companies": sum(1 for item in companies if len(item["contacts"]) >= 2),
        "blockers": len(blockers),
        "warnings": len(warnings),
        "hints": len(hints),
    }
    return {
        "summary": summary,
        "blockers": blockers,
        "warnings": warnings,
        "hints": hints,
        "companies": companies,
        "can_confirm": not blockers,
    }


def place_company(signal_level: str, destination: str, unrecognized: bool) -> str:
    if unrecognized or destination == "base":
        return "base"
    if destination == "pool":
        if signal_level == "confirm":
            return "pool"
        if signal_level == "leading":
            return "incubation"
        return "base"
    if signal_level == "confirm":
        return "reachable"
    if signal_level == "leading":
        return "incubation"
    return "trial"


def _blocked_file(message: str) -> dict[str, Any]:
    blocker = {"code": "B5", "excel_row": "", "customer_code": "", "company": "", "message": message}
    return {
        "summary": {
            "excel_rows": 0,
            "companies": 0,
            "contacts": 0,
            "multi_contact_companies": 0,
            "blockers": 1,
            "warnings": 0,
            "hints": 0,
        },
        "blockers": [blocker],
        "warnings": [],
        "hints": [],
        "companies": [],
        "can_confirm": False,
    }


def _read_sheet(content: bytes) -> tuple[list[tuple[int, list[str]]], list[str]]:
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = book.worksheets[0]
    table: list[tuple[int, list[str]]] = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        cells = [_cell(value) for value in row]
        if any(cells):
            table.append((index, cells))
    book.close()
    if not table:
        raise ValueError("表格是空的")
    header_row, headers = table[0]
    del header_row
    return table[1:], headers


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _header_map(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        title = header.replace(" ", "")
        if not title:
            continue
        for field, aliases in HEADER_FIELDS.items():
            if field in mapping:
                continue
            if any(alias.replace(" ", "") == title or alias.replace(" ", "") in title for alias in aliases):
                mapping[field] = index
    return mapping


def _parse_row(raw: list[str], mapping: dict[str, int], excel_row: int) -> dict[str, Any]:
    def take(field: str) -> str:
        index = mapping.get(field)
        if index is None or index >= len(raw):
            return ""
        return raw[index]

    item = {
        "excel_row": excel_row,
        "customer_code": take("customer_code"),
        "company": take("company"),
        "contact": take("contact"),
        "mobile_raw": take("mobile"),
        "landline_raw": take("landline"),
        "role": take("role"),
        "product": take("product"),
        "sales": take("sales"),
        "department": take("department"),
        "coach": take("coach"),
        "email": take("email"),
        "province": take("province"),
        "city": take("city"),
        "signals": take("signals"),
        "notes": take("notes"),
    }
    item["has_content"] = any(item[key] for key in ("customer_code", "company", "contact", "mobile_raw", "landline_raw", "product"))
    return item


def _group(parsed: list[dict[str, Any]], blockers: list[dict[str, str]]) -> dict[str, list[dict[str, Any]]]:
    names: dict[str, set[str]] = defaultdict(set)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in parsed:
        code = item["customer_code"]
        if not code or not item["company"]:
            continue
        names[code].add(item["company"])
        grouped[code].append(item)
    for code, seen in names.items():
        if len(seen) > 1:
            sample = grouped[code][0]
            blockers.append(_issue("B3", sample, "编码 %s 对应两个名称，请先统一" % code))
            grouped.pop(code, None)
    ready: dict[str, list[dict[str, Any]]] = {}
    for code, items in grouped.items():
        if not _any_number(items):
            blockers.append(_issue("B4", items[0], "%s 没有任何号码，无法进入工作台" % items[0]["company"]))
            continue
        ready[code] = items
    return ready


def _any_number(items: list[dict[str, Any]]) -> bool:
    return any(item["mobile_raw"] or item["landline_raw"] for item in items)


def _company_card(
    code: str,
    items: list[dict[str, Any]],
    warnings: list[dict[str, str]],
    hints: list[dict[str, str]],
) -> dict[str, Any]:
    company = items[0]["company"]
    products = _product_union(items, code, company, warnings)
    translated = catalog.translate_deals(products, _extra_rules())
    for original in translated["unrecognized"]:
        warnings.append(
            {
                "code": "W5",
                "excel_row": "",
                "customer_code": code,
                "company": company,
                "message": "成交原文未命中对照：%s。该公司只能进公海" % original,
            }
        )
    contacts, contact_warnings, contact_hints = _contacts(code, company, items)
    warnings.extend(contact_warnings)
    hints.extend(contact_hints)
    if not any(item["province"] or item["city"] for item in items):
        hints.append(_issue("I6", items[0], "%s 缺少省份或城市" % company))
    if any(not item["email"] for item in items):
        hints.append(_issue("I5", items[0], "%s 有联系人缺少邮箱或座机以外的通道，按实际号码展示" % company))
    classified = catalog.classify_company(company)
    signals = "\n".join(item["signals"] for item in items if item["signals"])
    notes = "\n".join(item["notes"] for item in items if item["notes"])
    level, leading, confirm = catalog.business_signals(classified["industry_code"], "\n".join((signals, notes)))
    preferred = next((item for item in contacts if item.get("preferred")), contacts[0] if contacts else None)
    return {
        "customer_code": code,
        "company": company,
        "purchased_display": "、".join(products),
        "scopes": translated["scopes"],
        "kmi_account": translated["kmi_account"],
        "consult_unspecified": translated["consult_unspecified"],
        "unrecognized": translated["unrecognized"],
        "industry_code": classified["industry_code"],
        "industry_name": classified["industry_name"],
        "industry_tags": classified["industry_tags"],
        "industry_status": classified["industry_status"],
        "entity_exception": classified["entity_exception"],
        "contacts": contacts,
        "preferred_name": (preferred or {}).get("name") or "",
        "signals_text": signals,
        "notes": notes,
        "signal_level": level,
        "leading_hits": leading,
        "confirm_hits": confirm,
        "related": [],
        "province": next((item["province"] for item in items if item["province"]), ""),
        "city": next((item["city"] for item in items if item["city"]), ""),
    }


def _product_union(
    items: list[dict[str, Any]],
    code: str,
    company: str,
    warnings: list[dict[str, str]],
) -> list[str]:
    found: list[str] = []
    filled = [item for item in items if item["product"]]
    if filled and len(filled) < len(items):
        warnings.append(
            {
                "code": "W4",
                "excel_row": "",
                "customer_code": code,
                "company": company,
                "message": "已购（展示）提升到公司级，原文：%s" % "、".join(item["product"] for item in filled),
            }
        )
    for item in filled:
        for part in _split_products(item["product"]):
            if part not in found:
                found.append(part)
    return found


def _split_products(text: str) -> list[str]:
    parts = [part.strip() for part in _SPLIT.split(text) if part.strip()]
    return parts or ([text.strip()] if text.strip() else [])


def _contacts(
    code: str,
    company: str,
    items: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, str]], list[dict[str, str]]]:
    warnings: list[dict[str, str]] = []
    hints: list[dict[str, str]] = []
    groups: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for item in items:
        name = item["contact"].strip()
        key = name or "row-%s" % item["excel_row"]
        channels = _channels(item, warnings, code, company)
        if key not in groups:
            groups[key] = {
                "id": new_id("CT"),
                "name": name,
                "role": item["role"],
                "roles_seen": [item["role"]] if item["role"] else [],
                "channels": [],
                "seen_numbers": set(),
                "rows": 1,
            }
            order.append(key)
        else:
            groups[key]["rows"] += 1
            if item["role"] and item["role"] not in groups[key]["roles_seen"]:
                groups[key]["roles_seen"].append(item["role"])
        for channel in channels:
            number = channel["number"]
            if number in groups[key]["seen_numbers"]:
                continue
            groups[key]["seen_numbers"].add(number)
            groups[key]["channels"].append(channel)
    contacts: list[dict[str, Any]] = []
    for key in order:
        group = groups[key]
        if group["rows"] > 1 and len(group["channels"]) == 1:
            warnings.append(
                {
                    "code": "W1",
                    "excel_row": "",
                    "customer_code": code,
                    "company": company,
                    "message": "%s 重复 %s 行，将合并为 1 个联系人" % (group["name"] or "未填姓名", group["rows"]),
                }
            )
        if group["rows"] > 1 and len(group["channels"]) > 1:
            warnings.append(
                {
                    "code": "W2",
                    "excel_row": "",
                    "customer_code": code,
                    "company": company,
                    "message": "%s：%s 个号码，将收成 1 人" % (group["name"] or "未填姓名", len(group["channels"])),
                }
            )
        role = _pick_role(group["roles_seen"])
        if len(group["roles_seen"]) > 1:
            hints.append(
                {
                    "code": "I7",
                    "excel_row": "",
                    "customer_code": code,
                    "company": company,
                    "message": "%s 职务不一致，已取：%s" % (group["name"], role),
                }
            )
        weak = _weak_role(role)
        if not role:
            hints.append(_hint(code, company, "I1", "%s 职务为空，排序靠后" % (group["name"] or "未填姓名")))
        elif weak:
            hints.append(_hint(code, company, "I2", "%s 职务「%s」过宽，排序低于主任、HRD、培训经理" % (group["name"], role)))
        if any(channel["switchboard"] for channel in group["channels"]) or "总机" in (group["name"] or ""):
            hints.append(_hint(code, company, "I3", "%s 有疑似总机号码，不作为一键首选" % (group["name"] or "未填姓名")))
        contacts.append(
            {
                "id": group["id"],
                "name": group["name"],
                "role": role,
                "weak_role": weak,
                "channels": group["channels"],
            }
        )
    if len(contacts) >= 2:
        hints.append(_hint(code, company, "I4", "%s 去重后 %s 位联系人，公司卡只露出首选" % (company, len(contacts))))
    _mark_preferred(contacts)
    return contacts, warnings, hints


def _channels(item: dict[str, Any], warnings: list[dict[str, str]], code: str, company: str) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    mobile_raw = item["mobile_raw"]
    if mobile_raw:
        kind = "landline" if _looks_landline(mobile_raw) else "mobile"
        if kind == "landline":
            warnings.append(
                {
                    "code": "W3",
                    "excel_row": str(item["excel_row"]),
                    "customer_code": code,
                    "company": company,
                    "message": "第 %s 行号码将改标为座机：%s" % (item["excel_row"], mobile_raw),
                }
            )
        found.append(_channel(kind, mobile_raw, item["contact"]))
    if item["landline_raw"]:
        found.append(_channel("landline", item["landline_raw"], item["contact"]))
    return found


def _channel(kind: str, number: str, contact: str) -> dict[str, Any]:
    switchboard = "总机" in contact or (kind == "landline" and _local_len(number) in {7, 8})
    return {"kind": kind, "number": number.strip(), "switchboard": switchboard, "invalid": False}


def _local_len(number: str) -> int:
    text = number.replace("—", "-")
    if "-" in text:
        return len(re.sub(r"\D", "", text.split("-")[-1]))
    digits = re.sub(r"\D", "", text)
    if digits.startswith("0") and len(digits) > 8:
        return len(digits) - (4 if len(digits) >= 12 else 3)
    return len(digits)


def _looks_landline(number: str) -> bool:
    text = number.strip()
    digits = re.sub(r"\D", "", text)
    if "-" in text or "—" in text:
        return True
    if text.startswith("0") or digits.startswith("0"):
        return True
    return False


def _pick_role(roles: list[str]) -> str:
    if not roles:
        return ""
    return sorted(roles, key=lambda role: (_role_rank(role), len(role)), reverse=True)[0]


def _role_rank(role: str) -> int:
    if not role.strip():
        return 0
    if _weak_role(role):
        return 1
    if any(token in role for token in ("主任", "HRD", "hrd", "总监", "总经理", "副总")):
        return 5
    if "经理" in role:
        return 3
    return 2


def _weak_role(role: str) -> bool:
    return role.strip().casefold() in WEAK_ROLES


def _mark_preferred(contacts: list[dict[str, Any]]) -> None:
    def score(contact: dict[str, Any]) -> tuple[int, int]:
        mobile = any(channel["kind"] == "mobile" and not channel["switchboard"] and not channel["invalid"] for channel in contact["channels"])
        return (1 if mobile else 0, _role_rank(contact.get("role") or ""))

    if not contacts:
        return
    best = max(range(len(contacts)), key=lambda index: score(contacts[index]))
    for index, contact in enumerate(contacts):
        contact["preferred"] = index == best


def _cross_contacts(companies: list[dict[str, Any]], warnings: list[dict[str, str]]) -> None:
    seen: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for company in companies:
        for contact in company["contacts"]:
            name = contact.get("name") or ""
            for channel in contact["channels"]:
                if channel["kind"] != "mobile":
                    continue
                digits = re.sub(r"\D", "", channel["number"])
                if name and digits:
                    seen[(name, digits)].append(company)
    reported: set[tuple[str, str]] = set()
    for (name, digits), owners in seen.items():
        codes = []
        for owner in owners:
            if owner["customer_code"] not in codes:
                codes.append(owner["customer_code"])
        if len(codes) < 2:
            continue
        key = (name, digits)
        if key in reported:
            continue
        reported.add(key)
        labels = "、".join(owner["company"] for owner in owners if owner["customer_code"] in codes)
        warnings.append(
            {
                "code": "W6",
                "excel_row": "",
                "customer_code": codes[0],
                "company": labels,
                "message": "%s 家公司共用联系人 %s，建议标关联" % (len(codes), name),
            }
        )
        for owner in owners:
            for other in owners:
                if other["customer_code"] == owner["customer_code"]:
                    continue
                pair = {"customer_code": other["customer_code"], "company": other["company"]}
                if pair not in owner["related"]:
                    owner["related"].append(pair)


def _link_known_pairs(companies: list[dict[str, Any]]) -> None:
    names = [item["company"] for item in companies]
    if any("无限家" in name for name in names) and any("十月知行" in name for name in names):
        for item in companies:
            if "无限家" in item["company"] and "关联十月知行" not in item["industry_tags"]:
                item["industry_tags"].append("关联十月知行")
            if "十月知行" in item["company"] and "关联无限家" not in item["industry_tags"]:
                item["industry_tags"].append("关联无限家")


def _sales_columns(parsed: list[dict[str, Any]], warnings: list[dict[str, str]]) -> None:
    values = []
    for item in parsed:
        values.extend(item[field] for field in ("sales", "department", "coach") if item[field])
    if not values:
        return
    if all(value.strip() in {"", "公海"} for value in values):
        warnings.append(
            {
                "code": "W7",
                "excel_row": "",
                "customer_code": "",
                "company": "",
                "message": "将忽略销售、部门、军师这三列，客户进入公海，由管理员再分配",
            }
        )


def _extra_rules() -> list[dict]:
    data = read_object("deal_map", {})
    rows = data.get("rows") if isinstance(data, dict) else []
    return rows if isinstance(rows, list) else []


def _issue(code: str, item: dict[str, Any], message: str) -> dict[str, str]:
    return {
        "code": code,
        "excel_row": str(item.get("excel_row") or ""),
        "customer_code": item.get("customer_code") or "",
        "company": item.get("company") or "",
        "message": message,
    }


def _hint(code: str, company: str, issue: str, message: str) -> dict[str, str]:
    return {"code": issue, "excel_row": "", "customer_code": code, "company": company, "message": message}
